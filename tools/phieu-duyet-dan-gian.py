# -*- coding: utf-8 -*-
"""Dung phieu duyet cum dan gian cho bac si: docs/dan-gian/phieu-duyet-cum-tu.docx.
Doc thang tu `src.dan_gian` de phieu va ma khong lech nhau. Chay: python tools/phieu-duyet-dan-gian.py"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Mm, Pt  # noqa: E402

from src import dan_gian  # noqa: E402

RA = Path(__file__).resolve().parents[1] / "docs" / "dan-gian" / "phieu-duyet-cum-tu.docx"
PHONG = "Times New Roman"

NGHIA = {  # nghia theo tu dien, viet lai ngan (khong chep nguyen van)
    "tháo bụng": "tiêu chảy (cách nói lịch sự)",
    "bao tử": "dạ dày",
    "xửng vửng": "choáng váng, hoa mắt",
    "ói": "nôn",
    "thở nghẹt": "(từ \"nghẹt\") có cảm giác khó thở",
    "bần thần": "rã rượi, mệt mỏi, suy nhược",
    "nhức bưng óc": "nhức đầu dữ dội",
    "xót xáy": "ngứa ngáy, đau rát nhẹ",
    "khậm khạc": "ho khạc dai dẳng",
    "ốm": "Nam Bộ: gầy (\"cà vom: ốm mà cao\")",
    "lên ban đỏ": "nổi đỏ cả mình, có thể là sởi hoặc sốt xuất huyết",
    "đau lộn ruột": "đau cả bụng (nhưng cũng nói \"cười lộn ruột\")",
}
LOAI = {dan_gian.TUONG_DUONG: "Rõ nghĩa", dan_gian.MO_HO: "Mơ hồ", dan_gian.TINH_CHAT: "Tả tính chất"}


def _p(r, co=11, dam=False, nghieng=False):
    r.font.name, r.font.size, r.font.bold, r.font.italic = PHONG, Pt(co), dam, nghieng
    rf = r._element.get_or_add_rPr().find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        r._element.get_or_add_rPr().append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(k), PHONG)


def doan(doc, van, co=11, dam=False, nghieng=False, sau=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(sau)
    _p(p.add_run(van), co, dam, nghieng)
    return p


def main():
    doc = Document()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Mm(297), Mm(210)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, m, Mm(14))
    doan(doc, "Phiếu nhờ bác sĩ duyệt: cách người bệnh nói về triệu chứng", 14, True, sau=2)
    doan(doc, "Dự án dùng các cách nói dưới đây để thử xem máy có hiểu nhầm lời bệnh nhân không. "
              "Các cụm lấy từ Từ điển Từ ngữ Nam Bộ (Huỳnh Công Tín, 2007). Nhờ bác đánh dấu ✓ vào ô phù hợp; "
              "chỗ nào nghĩa sai xin bác ghi nghĩa đúng. Khoảng 10–15 phút.", 11, sau=2)
    doan(doc, "Rõ nghĩa = ghi thẳng thành thuật ngữ được.   Mơ hồ = phải hỏi lại mới biết.   "
              "Tả tính chất = nói cảm giác/mức độ, không phải tên bệnh.", 10, nghieng=True, sau=6)

    hang = [(m.dan_gian, m.chuan, NGHIA.get(m.dan_gian, ""), LOAI[m.loai]) for m in dan_gian.BANG
            if m.chuan != "đi lỏng"]
    tieu = ["#", "Bệnh nhân nói", "Thay cho", "Nghĩa theo từ điển", "Máy xếp",
            "Nghĩa đúng?", "Bác xếp: Rõ / Mơ hồ / Tính chất", "Ghi chú của bác"]
    bang = doc.add_table(rows=1 + len(hang), cols=len(tieu))
    bang.style = "Table Grid"
    rong = [8, 32, 28, 70, 24, 26, 42, 50]
    for j, t in enumerate(tieu):
        c = bang.cell(0, j)
        c.width = Mm(rong[j])
        _p(c.paragraphs[0].add_run(t), 10, True)
    for i, (dg, chuan, nghia, loai) in enumerate(hang, 1):
        vals = [str(i), f"“{dg}”", chuan, nghia, loai, "□ Đúng   □ Sai", "□ Rõ  □ Mơ hồ  □ Tính chất", ""]
        for j, v in enumerate(vals):
            c = bang.cell(i, j)
            c.width = Mm(rong[j])
            _p(c.paragraphs[0].add_run(v), 10, j == 1)
    doan(doc, "", sau=2)
    doan(doc, "Hai câu người bệnh hay nói thêm (máy không được ghi thành chẩn đoán):", 11, True, sau=2)
    for cau, _, _ in dan_gian.CHEN:
        doan(doc, f"• “{cau}”   Bác thấy nên ghi vào hồ sơ không?  □ Không ghi  "
                  f"□ Ghi là lời người bệnh kể  □ Khác: ……………………", 10, sau=2)
    doan(doc, "", sau=2)
    doan(doc, "Ba câu hỏi ngắn:", 11, True, sau=2)
    for q in ["1. Khi bệnh nhân miền Bắc nói “cháu ốm” và bệnh nhân miền Nam nói “cháu ốm quá”, "
              "bác hỏi lại thế nào để biết là gầy hay đang bệnh?",
              "2. “Lên ban đỏ”: bác có tự nghĩ ngay tới một bệnh nào không, hay luôn hỏi thêm?",
              "3. Bác ghi giúp 3–5 cách nói về triệu chứng mà bệnh nhân của bác hay dùng và dễ hiểu nhầm: "
              "……………………………………………………………………………………………………………………………"]:
        doan(doc, q, 10, sau=3)
    RA.parent.mkdir(parents=True, exist_ok=True)
    doc.save(RA)
    print("da ghi", RA, "-", len(hang), "cum")


if __name__ == "__main__":
    main()
